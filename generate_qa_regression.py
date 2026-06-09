from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')

def create_doc():
    doc = Document()
    
    title = doc.add_heading('BÁO CÁO REGRESSION TEST VÀ CROSS-BROWSER MATRIX', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, '1. Tóm tắt chu trình Regression Test (Kiểm thử Hồi quy)', 1)
    doc.add_paragraph('Bối cảnh: Sau khi team QA phát hiện 10 lỗi (Bug Report trước đó) trên phiên bản V1.0, team Dev đã tiến hành sửa lỗi và phát hành phiên bản V1.1.\n'
                      'Mục tiêu: Đảm bảo các lỗi cũ đã được khắc phục hoàn toàn (Re-test) và những thay đổi code không làm hỏng các tính năng cốt lõi đang hoạt động tốt (Regression Test).\n'
                      'Phạm vi hồi quy (Scope): Tập trung vào module Giỏ Hàng (Cart), Thanh Toán (Checkout) và Đăng Ký (Register).')
                      
    add_heading(doc, '2. Báo cáo Test Execution (Thực thi Kiểm thử)', 1)
    doc.add_paragraph('Dưới đây là Báo cáo Thực thi Kiểm thử tóm tắt kết quả chạy Hồi quy cho đợt Release V1.1:')
    
    # Bảng Test Execution
    exec_table = doc.add_table(rows=1, cols=2)
    exec_table.style = 'Table Grid'
    hdr = exec_table.rows[0].cells
    hdr[0].text = 'Hạng mục'
    hdr[1].text = 'Thông tin chi tiết'
    
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs: r.font.bold = True
            
    data = [
        ['Dự án', 'Shoes Store E-commerce'],
        ['Phiên bản (Build/Release)', 'v1.1.0-RC1'],
        ['Môi trường Test', 'Staging (Windows 11 / Chrome)'],
        ['Tổng số Test Case thực thi', '50 TC'],
        ['Số lượng TC Pass', '47 TC'],
        ['Số lượng TC Fail', '3 TC (Lỗi UI nhỏ)'],
        ['Số lượng TC Blocked', '0 TC'],
        ['Tỷ lệ Pass Rate', '94%'],
        ['Kết luận', 'Hệ thống ổn định. Cho phép Go-live (Release) lên Production.']
    ]
    
    for d in data:
        row = exec_table.add_row().cells
        row[0].text = d[0]
        row[1].text = d[1]
    
    doc.add_paragraph('\nNhận xét chung: Các lỗi nghiêm trọng (Critical/High) như "Thêm số lượng âm vào giỏ hàng" và "Đặt hàng không cần địa chỉ" đã được fix thành công. 3 TC Fail còn lại chỉ là lỗi hiển thị text (Low Priority), sẽ được đưa vào backlog để fix ở Sprint sau.')

    doc.add_page_break()

    add_heading(doc, '3. BT 4.3: Ma trận Test Đa Trình Duyệt (Cross-Browser Matrix)', 1)
    doc.add_paragraph('Chiến lược Cross-browser testing giúp đảm bảo web Shoes Store hiển thị tốt trên các nền tảng phổ biến nhất của người dùng cuối. Dưới đây là Ma trận kiểm thử hỗ trợ trình duyệt:')
    
    # Bảng Matrix
    matrix_table = doc.add_table(rows=1, cols=5)
    matrix_table.style = 'Table Grid'
    hdr = matrix_table.rows[0].cells
    hdr[0].text = 'Hệ điều hành (OS)'
    hdr[1].text = 'Google Chrome'
    hdr[2].text = 'Mozilla Firefox'
    hdr[3].text = 'Microsoft Edge'
    hdr[4].text = 'Apple Safari'
    
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs: r.font.bold = True
            
    matrix_data = [
        ['Windows 11 (Desktop)', 'Có (P1)', 'Có (P2)', 'Có (P2)', 'Không hỗ trợ'],
        ['macOS (Desktop)', 'Có (P1)', 'Có (P2)', 'Không Test', 'Có (P1)'],
        ['Android 13+ (Mobile)', 'Có (P1)', 'Có (P3)', 'Không Test', 'Không hỗ trợ'],
        ['iOS 16+ (Mobile)', 'Có (P1)', 'Không Test', 'Không Test', 'Có (P1)']
    ]
    
    for d in matrix_data:
        row = matrix_table.add_row().cells
        for i in range(5):
            row[i].text = d[i]
            
    doc.add_paragraph('\nChú thích các mức độ ưu tiên test (Priority):\n'
                      '- P1 (Priority 1): Trình duyệt chính, nhiều user dùng nhất. Phải test 100% Test Cases (Full Regression).\n'
                      '- P2 (Priority 2): Trình duyệt phụ. Chỉ test các luồng cơ bản (Smoke Test) và kiểm tra giao diện (UI).\n'
                      '- P3 (Priority 3): Ít user dùng. Chỉ test nhanh trên thiết bị thật (Sanity Test) nếu có thời gian.\n'
                      '- Không hỗ trợ / Không Test: Nền tảng không tương thích hoặc số lượng user dùng quá thấp (< 1%), không đáng để đầu tư effort test.')

    doc.save('QA_Regression_CrossBrowser.docx')

if __name__ == "__main__":
    create_doc()
    print("File QA_Regression_CrossBrowser.docx created successfully!")
