from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')

def create_doc():
    doc = Document()
    
    title = doc.add_heading('BÁO CÁO LỖI (BUG REPORT) VÀ MA TRẬN YÊU CẦU (RTM)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Jira Guide
    add_heading(doc, 'Hướng dẫn Dùng Jira/TestRail tạo bug', 1)
    doc.add_paragraph('Để hoàn thành yêu cầu "Sử dụng công cụ chuyên nghiệp" của Giảng viên, bạn cần làm theo các bước sau trên Jira:\n'
                      '1. Đăng nhập Jira, tạo một Project mới chọn template là "Scrum" hoặc "Kanban", đặt tên là "Shoes Store QA".\n'
                      '2. Ở góc trên cùng, bấm nút "Create" (Tạo mới).\n'
                      '3. Chọn Issue Type là "Bug" (Lỗi).\n'
                      '4. Ở trường Summary, copy dán Tiêu đề lỗi (Bug Summary) từ danh sách bên dưới.\n'
                      '5. Ở trường Description, điền theo format: Environment, Steps to reproduce, Expected Result, Actual Result.\n'
                      '6. Chọn mức độ cho Priority (Highest, High, Medium, Low) dựa theo bảng Bug Triage bên dưới.\n'
                      '7. Bấm "Create" và chụp ảnh màn hình cái thẻ Bug đó trên Jira dán vào báo cáo nếu cần.')

    doc.add_page_break()

    # 2. Bug Report & Bug Triage
    add_heading(doc, 'BT 5.1 & BT 5.2: Danh sách 10 Bug và Đánh giá Mức độ (Bug Triage)', 1)
    doc.add_paragraph('Bảng dưới đây liệt kê 10 lỗi được mô phỏng tìm thấy trên hệ thống Shoes Store, kết hợp với phần đánh giá Mức độ nghiêm trọng (Severity) và Mức độ ưu tiên (Priority).')
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Bug ID'
    hdr[1].text = 'Tiêu đề lỗi (Bug Summary)'
    hdr[2].text = 'Severity (Nghiêm trọng)'
    hdr[3].text = 'Priority (Ưu tiên)'
    
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs: r.font.bold = True

    bugs = [
        ['BUG-01', 'Có thể thêm số lượng âm (-5) vào giỏ hàng', 'Critical', 'High'],
        ['BUG-02', 'Tổng tiền (Total) không cập nhật khi xóa sản phẩm khỏi giỏ', 'High', 'High'],
        ['BUG-03', 'Đăng ký thành công dù Email nhập sai định dạng (vượt qua HTML5 validation)', 'Medium', 'Medium'],
        ['BUG-04', 'Nút "Thêm vào giỏ hàng" bị che khuất bởi footer trên giao diện Mobile', 'High', 'Medium'],
        ['BUG-05', 'Popup thông báo thêm giỏ hàng thành công hiển thị chữ tiếng Anh ("Success") thay vì tiếng Việt', 'Low', 'Low'],
        ['BUG-06', 'Hình ảnh sản phẩm bị méo (không giữ tỷ lệ) ở màn hình Chi tiết sản phẩm', 'Low', 'Low'],
        ['BUG-07', 'Cho phép đặt hàng thành công dù không nhập Địa chỉ giao hàng', 'Critical', 'Highest'],
        ['BUG-08', 'Crash ứng dụng (Lỗi 500) khi tìm kiếm bằng từ khóa có chứa ký tự đặc biệt (!@#)', 'Critical', 'High'],
        ['BUG-09', 'Số lượng hiển thị trên icon Giỏ hàng không tự cập nhật sau khi thêm sản phẩm (Phải F5)', 'Medium', 'Medium'],
        ['BUG-10', 'Tên sản phẩm quá dài bị tràn chữ, đè lên giá tiền trong trang Danh sách', 'Medium', 'Low']
    ]

    for b in bugs:
        row = table.add_row().cells
        row[0].text = b[0]
        row[1].text = b[1]
        row[2].text = b[2]
        row[3].text = b[3]

    doc.add_paragraph('\nLưu ý (Bug Triage):')
    doc.add_paragraph('- BUG-07 (Không nhập địa chỉ vẫn mua được) là lỗi Highest Priority vì ảnh hưởng trực tiếp đến quy trình giao hàng và doanh thu.')
    doc.add_paragraph('- BUG-01, BUG-02, BUG-08 là High Priority vì làm hỏng luồng thanh toán và gây lỗi server.')
    doc.add_paragraph('- BUG-05, BUG-06 là Low Priority vì chỉ ảnh hưởng nhỏ đến UI/UX, không chặn quy trình mua hàng.')

    doc.add_page_break()

    # 3. RTM
    add_heading(doc, 'BT 1.3: Ma trận truy xuất yêu cầu (Requirements Traceability Matrix - RTM)', 1)
    doc.add_paragraph('Bảng RTM giúp theo dõi mức độ bao phủ của Test Case đối với các Yêu cầu (Requirements), đồng thời liên kết trực tiếp với các lỗi (Bugs) được phát hiện.')

    rtm_table = doc.add_table(rows=1, cols=5)
    rtm_table.style = 'Table Grid'
    rtm_hdr = rtm_table.rows[0].cells
    rtm_hdr[0].text = 'Req ID'
    rtm_hdr[1].text = 'Mô tả Yêu cầu (Requirement)'
    rtm_hdr[2].text = 'Test Case ID Coverage'
    rtm_hdr[3].text = 'Trạng thái'
    rtm_hdr[4].text = 'Bug ID Liên kết'
    
    for cell in rtm_hdr:
        for p in cell.paragraphs:
            for r in p.runs: r.font.bold = True
            
    rtm_data = [
        ['REQ-01', 'Hệ thống cho phép người dùng đăng ký tài khoản mới', 'TC_REG_01 -> TC_REG_20', 'Failed', 'BUG-03'],
        ['REQ-02', 'Hệ thống cho phép tìm kiếm sản phẩm theo tên', 'TC_SEARCH_01 -> TC_SEARCH_05', 'Failed', 'BUG-08'],
        ['REQ-03', 'Hệ thống hiển thị đúng hình ảnh và thông tin chi tiết sản phẩm', 'TC_PROD_01 -> TC_PROD_05', 'Failed', 'BUG-06, BUG-10'],
        ['REQ-04', 'Khách hàng có thể thêm sản phẩm vào giỏ hàng', 'TC_CART_01 -> TC_CART_08', 'Failed', 'BUG-01, BUG-04, BUG-09'],
        ['REQ-05', 'Hệ thống tính đúng tổng tiền và quản lý giỏ hàng', 'TC_CART_09 -> TC_CART_15', 'Failed', 'BUG-02'],
        ['REQ-06', 'Hệ thống yêu cầu nhập đầy đủ thông tin giao hàng khi Thanh toán', 'TC_CHECKOUT_01 -> TC_CHECKOUT_10', 'Failed', 'BUG-07']
    ]

    for d in rtm_data:
        row = rtm_table.add_row().cells
        for i in range(5):
            row[i].text = d[i]

    doc.add_paragraph('\nÝ nghĩa bảng RTM:\n'
                      '- Đảm bảo mọi Yêu cầu (REQ) đều có ít nhất một Test Case bao phủ.\n'
                      '- Đảm bảo mọi Bug tìm thấy đều có thể truy xuất ngược lại Yêu cầu nào đang bị hỏng.\n'
                      '- Giúp Project Manager (PM) nắm được tính năng nào đang gặp nhiều lỗi nhất (Vd: REQ-04).')

    doc.save('QA_BugReport_RTM.docx')

if __name__ == "__main__":
    create_doc()
    print("File QA_BugReport_RTM.docx created successfully!")
