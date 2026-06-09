import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')

def create_qa_doc():
    doc = Document()
    
    # Tiêu đề tài liệu
    title = doc.add_heading('BÁO CÁO THỰC HÀNH QC AUTOMATION - TUẦN 1', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Dự án: Hệ thống Thương mại Điện tử Shoes Store').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Người thực hiện: Nhóm/Cá nhân').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Ngày tạo: Hôm nay').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # 1. TEST PLAN
    add_heading(doc, 'PHẦN 1: TEST PLAN (KẾ HOẠCH KIỂM THỬ)', 1)
    
    add_heading(doc, '1.1 Giới thiệu (Introduction)', 2)
    doc.add_paragraph('Tài liệu này trình bày kế hoạch kiểm thử (Test Plan) cho website e-commerce Shoes Store. Mục tiêu là đảm bảo chất lượng cho các chức năng cốt lõi của hệ thống trước khi đưa vào tự động hóa (Automation) ở các tuần tiếp theo.')
    
    add_heading(doc, '1.2 Phạm vi kiểm thử (Test Scope)', 2)
    doc.add_paragraph('Các chức năng nằm trong phạm vi kiểm thử (In-scope):')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Quản lý tài khoản (Đăng ký, Đăng nhập, Quên mật khẩu).')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Quản lý hồ sơ (Cập nhật thông tin, Sổ địa chỉ).')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Tính năng sản phẩm (Tìm kiếm, Chi tiết sản phẩm, Wishlist, Đánh giá).')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Giỏ hàng và Thanh toán (Checkout).')

    add_heading(doc, '1.3 Môi trường kiểm thử (Test Environment)', 2)
    doc.add_paragraph('Môi trường cục bộ (Localhost) chạy trên Django Server.\nTrình duyệt: Google Chrome, Firefox.\nHệ điều hành: Windows/macOS.')
    
    add_heading(doc, '1.4 Tiêu chí đánh giá (Pass/Fail Criteria)', 2)
    doc.add_paragraph('- Pass: Kết quả thực tế (Actual Result) khớp hoàn toàn với kết quả mong đợi (Expected Result) và không sinh ra lỗi hệ thống (500 Server Error).\n- Fail: Kết quả thực tế sai lệch với mong đợi hoặc ứng dụng bị crash.')
    
    doc.add_page_break()
    
    # 2. TEST CASES
    add_heading(doc, 'PHẦN 2: TEST CASES (KỊCH BẢN KIỂM THỬ)', 1)
    doc.add_paragraph('Dưới đây là danh sách các Test Cases tiêu biểu cho các module cốt lõi của ứng dụng.')
    
    # Define test cases data
    test_cases = [
        # Authentication
        ["TC01_AUTH_01", "Đăng ký tài khoản thành công với thông tin hợp lệ", "Người dùng đang ở trang /register/", "1. Nhập Username\n2. Nhập Email\n3. Nhập Password\n4. Nhập Confirm Password\n5. Click 'Đăng ký'", "Hệ thống thông báo đăng ký thành công và chuyển hướng đến trang Đăng nhập."],
        ["TC01_AUTH_02", "Đăng ký thất bại do mật khẩu không khớp", "Người dùng đang ở trang /register/", "1. Nhập thông tin hợp lệ\n2. Cố tình nhập Password và Confirm Password khác nhau\n3. Click 'Đăng ký'", "Hệ thống báo lỗi 'Mật khẩu không khớp' và không tạo tài khoản."],
        ["TC01_AUTH_03", "Đăng nhập thành công với tài khoản đúng", "Có sẵn tài khoản hợp lệ. Ở trang /login/", "1. Nhập Username hợp lệ\n2. Nhập Password hợp lệ\n3. Click 'Đăng nhập'", "Đăng nhập thành công, chuyển hướng về Trang chủ và hiển thị tên người dùng trên Navbar."],
        ["TC01_AUTH_04", "Đăng nhập thất bại do sai mật khẩu", "Có sẵn tài khoản hợp lệ. Ở trang /login/", "1. Nhập Username hợp lệ\n2. Nhập sai Password\n3. Click 'Đăng nhập'", "Hệ thống hiển thị lỗi 'Tài khoản hoặc mật khẩu không chính xác'."],
        
        # Dashboard & Address
        ["TC02_USER_01", "Cập nhật thông tin cá nhân thành công", "Đã đăng nhập. Ở trang /profile/ (Tab Hồ sơ)", "1. Sửa First name, Last name\n2. Click 'Lưu thay đổi'", "Thông tin được cập nhật vào database. Hiển thị thông báo thành công."],
        ["TC02_USER_02", "Thêm địa chỉ giao hàng mới", "Đã đăng nhập. Ở trang /profile/ (Tab Sổ địa chỉ)", "1. Click 'Thêm địa chỉ mới'\n2. Nhập đủ Họ tên, SĐT, Địa chỉ\n3. Click 'Lưu'", "Địa chỉ mới xuất hiện trong danh sách. Nếu là địa chỉ đầu tiên, nó mặc định là Default."],
        
        # Product & Cart
        ["TC03_PROD_01", "Xem chi tiết một sản phẩm", "Đang ở trang danh sách sản phẩm", "1. Click vào một sản phẩm bất kỳ", "Chuyển sang trang Product Detail. Hiển thị đúng tên, giá, hình ảnh, sizes và các đánh giá."],
        ["TC03_PROD_02", "Thêm sản phẩm vào Wishlist", "Đã đăng nhập. Đang ở trang Product Detail", "1. Click vào icon Trái tim (Wishlist)", "Icon chuyển sang màu đỏ. Hiện popup 'Đã thêm vào yêu thích'."],
        ["TC03_CART_01", "Thêm sản phẩm vào giỏ hàng thành công", "Đã đăng nhập. Ở trang Product Detail", "1. Chọn Size\n2. Nhập Số lượng (vd: 1)\n3. Click 'Thêm vào giỏ'", "Sản phẩm được thêm vào giỏ. Nút đổi thành 'ĐÃ THÊM!'. Số lượng trên Navbar tăng lên 1."],
        ["TC03_CART_02", "Thêm sản phẩm vào giỏ hàng khi chưa chọn Size", "Đã đăng nhập. Ở trang Product Detail", "1. Bỏ qua việc chọn Size\n2. Click 'Thêm vào giỏ'", "Hệ thống chặn lại và hiện thông báo 'Vui lòng chọn size trước khi thêm'."],
        ["TC03_CART_03", "Cập nhật số lượng trong giỏ hàng", "Đã đăng nhập. Ở trang /cart/", "1. Bấm nút '+' để tăng số lượng sản phẩm", "Tổng tiền của sản phẩm đó và tổng đơn hàng được cập nhật tương ứng."],
        
        # Checkout
        ["TC04_CHK_01", "Thanh toán thành công qua hình thức COD", "Đã đăng nhập. Có sản phẩm trong giỏ. Đang ở trang /checkout/", "1. Chọn địa chỉ từ Dropdown (tự động điền)\n2. Chọn COD\n3. Click 'Đặt hàng'", "Đơn hàng được tạo thành công với trạng thái Pending. Chuyển sang trang Cảm ơn."],
        ["TC04_CHK_02", "Kiểm tra Auto-fill địa chỉ từ Sổ địa chỉ", "Đã đăng nhập. Có lưu ít nhất 1 địa chỉ. Đang ở /checkout/", "1. Mở dropdown 'Chọn từ Sổ địa chỉ'\n2. Chọn 1 địa chỉ", "Các field Họ tên, Số điện thoại, Địa chỉ tự động được điền đúng dữ liệu của địa chỉ đã chọn."],
        ["TC04_CHK_03", "Báo lỗi khi đặt hàng thiếu thông tin", "Đã đăng nhập. Ở trang /checkout/", "1. Xóa rỗng trường Số điện thoại\n2. Click 'Đặt hàng'", "Form chặn submit và yêu cầu nhập trường Số điện thoại."],
        
        # Reviews
        ["TC05_REV_01", "Chặn đánh giá nếu chưa mua hàng", "Đã đăng nhập. Ở trang Product Detail của SP chưa mua", "1. Kéo xuống phần Đánh giá", "Không hiển thị Form đánh giá. Hiển thị thông báo 'Bạn cần mua và nhận thành công... để đánh giá'."],
        ["TC05_REV_02", "Gửi đánh giá thành công (Verified Purchase)", "Đã đăng nhập. Đã có đơn hàng Completed với SP này.", "1. Kéo xuống phần Đánh giá\n2. Chọn 5 sao\n3. Nhập comment\n4. Upload ảnh\n5. Click 'Gửi đánh giá'", "Review hiển thị trên danh sách kèm badge màu xanh 'Đã mua hàng' và hình ảnh đính kèm."]
    ]
    
    # Tạo bảng
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Test Case ID'
    hdr_cells[1].text = 'Title (Mô tả)'
    hdr_cells[2].text = 'Pre-conditions (Tiền điều kiện)'
    hdr_cells[3].text = 'Test Steps (Các bước test)'
    hdr_cells[4].text = 'Expected Results (Kết quả mong đợi)'
    
    # Make header bold
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Thêm data
    for tc in test_cases:
        row_cells = table.add_row().cells
        row_cells[0].text = tc[0]
        row_cells[1].text = tc[1]
        row_cells[2].text = tc[2]
        row_cells[3].text = tc[3]
        row_cells[4].text = tc[4]
        
    doc.save('QA_Week1_TestPlan_TestCases.docx')

if __name__ == "__main__":
    create_qa_doc()
    print("Document successfully created: QA_Week1_TestPlan_TestCases.docx")
