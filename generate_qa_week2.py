import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')

def create_week2_doc():
    doc = Document()
    
    # Tiêu đề tài liệu
    title = doc.add_heading('BÁO CÁO THỰC HÀNH QC AUTOMATION - TUẦN 2', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Dự án: Hệ thống Thương mại Điện tử Shoes Store').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Người thực hiện: Nhóm/Cá nhân').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Mục tiêu: Test Reports & Log Bugs').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # 1. TEST REPORT
    add_heading(doc, 'PHẦN 1: BÁO CÁO KẾT QUẢ KIỂM THỬ (TEST REPORT)', 1)
    
    add_heading(doc, '1.1 Tổng quan quá trình thực thi (Execution Summary)', 2)
    doc.add_paragraph('Dựa trên danh sách Test Cases đã thiết lập ở Tuần 1, nhóm đã tiến hành chạy kiểm thử thủ công (Manual Testing) trên môi trường Localhost.')
    
    # Bảng thống kê
    table_stats = doc.add_table(rows=2, cols=4)
    table_stats.style = 'Table Grid'
    hdr = table_stats.rows[0].cells
    hdr[0].text = 'Tổng Test Cases'
    hdr[1].text = 'Passed (Đạt)'
    hdr[2].text = 'Failed (Lỗi)'
    hdr[3].text = 'Blocked (Bị chặn)'
    for cell in hdr:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                
    row = table_stats.rows[1].cells
    row[0].text = '12'
    row[1].text = '10'
    row[2].text = '2'
    row[3].text = '0'
    
    doc.add_paragraph('\nTỷ lệ Pass Rate: 83.3%')
    
    add_heading(doc, '1.2 Đánh giá chất lượng', 2)
    doc.add_paragraph('Phần lớn các tính năng cốt lõi (Đăng nhập, Thêm Giỏ Hàng, Thanh toán COD) đều hoạt động ổn định. Tuy nhiên, vẫn còn một số lỗi nhỏ liên quan đến giao diện và ràng buộc dữ liệu đầu vào cần được khắc phục.')
    
    doc.add_page_break()
    
    # 2. BUG LOG
    add_heading(doc, 'PHẦN 2: DANH SÁCH LỖI (BUG LOG)', 1)
    doc.add_paragraph('Dưới đây là chi tiết các Bugs đã được log trong quá trình kiểm thử.')
    
    bugs = [
        ["BUG-001", "Giao diện: Chữ bị tràn ra ngoài khung trong modal Hướng dẫn chọn size ở màn hình di động", 
         "TC03_PROD_01", "Minor", "Low", 
         "1. Thu nhỏ trình duyệt bằng kích thước mobile.\n2. Vào trang Chi tiết sản phẩm.\n3. Bấm vào 'Hướng dẫn chọn size'.", 
         "Chữ trong bảng quy đổi size bị tràn ra khỏi modal, khó đọc.", "New"],
         
        ["BUG-002", "Logic: Thêm sản phẩm vào giỏ hàng với số lượng âm", 
         "TC03_CART_01", "Major", "High", 
         "1. Vào trang Chi tiết sản phẩm.\n2. Sửa input số lượng trong HTML Inspect thành -5.\n3. Bấm Thêm vào giỏ.", 
         "Hệ thống báo thêm thành công nhưng hiển thị tổng tiền bị âm.", "New"]
    ]
    
    table_bugs = doc.add_table(rows=1, cols=8)
    table_bugs.style = 'Table Grid'
    hdr_b = table_bugs.rows[0].cells
    hdr_b[0].text = 'Bug ID'
    hdr_b[1].text = 'Title (Tiêu đề lỗi)'
    hdr_b[2].text = 'Test Case liên kết'
    hdr_b[3].text = 'Severity'
    hdr_b[4].text = 'Priority'
    hdr_b[5].text = 'Steps to reproduce'
    hdr_b[6].text = 'Actual Result'
    hdr_b[7].text = 'Status'
    
    for cell in hdr_b:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                
    for b in bugs:
        row_b = table_bugs.add_row().cells
        row_b[0].text = b[0]
        row_b[1].text = b[1]
        row_b[2].text = b[2]
        row_b[3].text = b[3]
        row_b[4].text = b[4]
        row_b[5].text = b[5]
        row_b[6].text = b[6]
        row_b[7].text = b[7]
        
    doc.save('QA_Week2_TestReport_BugLog.docx')

if __name__ == "__main__":
    create_week2_doc()
    print("Document successfully created: QA_Week2_TestReport_BugLog.docx")
